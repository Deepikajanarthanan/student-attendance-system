from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

import database
from attendance import (add_or_update_holiday, add_student, attendance_exists,
                        delete_attendance, get_absent_roll_nos, get_course_details,
                        get_register_session, get_session_class_dates, get_session_status,
                        get_students, import_attendance_by_roll, save_attendance,
                        save_course_details, save_register_session)
from bulk_import import parse_attendance_workbook
from reports import (attendance_matrix, calculate_attendance_marks, excel_export,
                     generate_register_excel, student_detail, summary_report, summary_report_between)


@pytest.fixture(autouse=True)
def temporary_database(tmp_path, monkeypatch):
    monkeypatch.setattr(database, "DB_PATH", Path(tmp_path) / "attendance.db")
    database.init_db()
    for reg, roll, name in [("101", "01", "Arun"), ("102", "02", "Bala"), ("103", "03", "Deepa"),
                            ("104", "04", "Hari"), ("105", "05", "Priya")]:
        add_student(reg, roll, name)


def test_absent_students_and_everyone_else_present():
    present, absent = save_attendance("2026-09-13", "02,04")
    assert (present, absent) == (3, 2)
    assert get_absent_roll_nos("2026-09-13") == ["02", "04"]
    summary = summary_report(2026, 9)
    assert summary.loc[summary["Reg No"] == "101", "Present"].iloc[0] == 1
    assert summary.loc[summary["Reg No"] == "102", "Absent"].iloc[0] == 1


def test_empty_absent_list_means_everyone_present():
    assert save_attendance("2026-09-14", "") == (5, 0)
    assert get_absent_roll_nos("2026-09-14") == []


def test_holiday_is_excluded_and_shown_in_matrix():
    save_attendance("2026-09-13", "02")
    add_or_update_holiday("2026-09-15", "College Holiday")
    summary = summary_report(2026, 9)
    assert summary["Working Days"].iloc[0] == 1
    matrix = attendance_matrix(2026, 9)
    assert set(matrix["15-Sep"]) == {"H"}


def test_invalid_reg_no_does_not_save():
    with pytest.raises(ValueError, match="999"):
        save_attendance("2026-09-13", "02,999,04")
    assert not attendance_exists("2026-09-13")


def test_duplicate_input_is_deduplicated_and_update_replaces_date():
    assert save_attendance("2026-09-13", "02,02,04") == (3, 2)
    # Replacing with a valid registration number recalculates all statuses.
    assert save_attendance("2026-09-13", "02,05", replace=True) == (3, 2)
    assert get_absent_roll_nos("2026-09-13") == ["02", "05"]


def test_delete_attendance_and_excel_export():
    save_attendance("2026-09-13", "02")
    add_or_update_holiday("2026-09-15", "College Holiday")
    assert excel_export(2026, 9).startswith(b"PK")
    delete_attendance("2026-09-13")
    assert not attendance_exists("2026-09-13")


def test_combined_months_report_uses_all_submitted_days():
    save_attendance("2026-09-13", "02")
    save_attendance("2026-10-01", "")
    report = summary_report_between("2026-09-01", "2026-11-30")
    arun = report.loc[report["Reg No"] == "101"].iloc[0]
    bala = report.loc[report["Reg No"] == "102"].iloc[0]
    assert (arun["Working Days"], arun["Present"], arun["Attendance %"]) == (2, 2, 100.0)
    assert (bala["Working Days"], bala["Absent"], bala["Attendance %"]) == (2, 1, 50.0)


def test_bulk_attendance_import_uses_roll_numbers_and_is_complete():
    imported = import_attendance_by_roll({
        "2026-09-13": {"01": "P", "02": "A", "03": "P", "04": "A", "05": "P"},
        "2026-09-14": {"01": "P", "02": "P", "03": "P", "04": "P", "05": "P"},
    })
    assert imported == 2
    assert get_absent_roll_nos("2026-09-13") == ["02", "04"]
    with pytest.raises(ValueError, match="missing attendance"):
        import_attendance_by_roll({"2026-09-15": {"01": "P"}})


def test_excel_sheet_parser_reads_roll_numbers_dates_and_pa_values():
    original = pd.DataFrame({
        "Roll No": ["01", "02"], "Reg No": ["101", "102"], "Name": ["Arun", "Bala"],
        "13-09-2026": ["P", "A"], "14-09-2026": ["P", "P"],
    })
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        original.to_excel(writer, index=False)
    parsed = parse_attendance_workbook(output.getvalue(), "Sheet1", date(2026, 9, 13), date(2026, 9, 14))
    assert parsed.records["2026-09-13"] == {"01": "P", "02": "A"}


def test_calculate_attendance_marks_scheme():
    # Rule: reduce 1 mark for every 3 classes absent (from max 5 marks)
    assert calculate_attendance_marks(0) == 5
    assert calculate_attendance_marks(1) == 5
    assert calculate_attendance_marks(2) == 5
    assert calculate_attendance_marks(3) == 4   # 1 mark reduced for 3 classes absent
    assert calculate_attendance_marks(4) == 4
    assert calculate_attendance_marks(5) == 4
    assert calculate_attendance_marks(6) == 3   # 2 marks reduced for 6 classes absent
    assert calculate_attendance_marks(8) == 3
    assert calculate_attendance_marks(9) == 2   # 3 marks reduced for 9 classes absent
    assert calculate_attendance_marks(11) == 2
    assert calculate_attendance_marks(12) == 1  # 4 marks reduced for 12 classes absent
    assert calculate_attendance_marks(14) == 1
    assert calculate_attendance_marks(15) == 0  # 5 marks reduced
    assert calculate_attendance_marks(30) == 0

    # Percentage conversions on 40-hour course
    assert calculate_attendance_marks(100.0, total_periods=40) == 5
    assert calculate_attendance_marks(95.0, total_periods=40) == 5   # 2 absent
    assert calculate_attendance_marks(92.5, total_periods=40) == 4   # 3 absent
    assert calculate_attendance_marks(85.0, total_periods=40) == 3   # 6 absent
    assert calculate_attendance_marks(77.5, total_periods=40) == 2   # 9 absent
    assert calculate_attendance_marks(70.0, total_periods=40) == 1   # 12 absent
    assert calculate_attendance_marks(50.0, total_periods=40) == 0   # 20 absent


def test_session_class_dates_and_status():
    # September 2026: 2026-09-01 is Tuesday, 2026-09-03 is Thursday, 2026-09-07 is Monday, 2026-09-10 is Thursday
    class_dates = get_session_class_dates(date(2026, 9, 1), date(2026, 9, 10), weekdays=[0, 3], exclude_holidays=True)
    assert class_dates == [date(2026, 9, 3), date(2026, 9, 7), date(2026, 9, 10)]

    status = get_session_status(class_dates)
    assert status["total_dates"] == 3
    assert status["is_complete"] is False
    assert status["next_unmarked_date"] == date(2026, 9, 3)

    # Save attendance for 2 dates
    save_attendance("2026-09-03", "02")
    save_attendance("2026-09-07", "")
    status = get_session_status(class_dates)
    assert status["is_complete"] is False
    assert status["next_unmarked_date"] == date(2026, 9, 10)

    # Save last date
    save_attendance("2026-09-10", "01,04")
    status = get_session_status(class_dates)
    assert status["is_complete"] is True
    assert status["next_unmarked_date"] is None


def test_generate_register_excel_structure_and_marks():
    import openpyxl
    class_dates = [date(2026, 9, 3), date(2026, 9, 7), date(2026, 9, 10)]
    save_attendance("2026-09-03", "02")
    save_attendance("2026-09-07", "")
    save_attendance("2026-09-10", "02")

    # 1 hour per day, Mon = 1st hour, Thu = 4th hour
    excel_bytes = generate_register_excel(date(2026, 9, 1), date(2026, 9, 10), class_dates=class_dates, periods_per_day=1, weekday_periods={0: 1, 3: 4})
    wb = openpyxl.load_workbook(BytesIO(excel_bytes))
    ws = wb.active
    assert ws.title == "Attendance Register"

    # Row 3: PERIODS header and weekday period hours (Thu=4, Mon=1, Thu=4)
    assert ws.cell(row=3, column=3).value == "PERIODS"
    assert ws.cell(row=3, column=4).value == 4
    assert ws.cell(row=3, column=5).value == 1
    assert ws.cell(row=3, column=6).value == 4
    assert ws.cell(row=3, column=7).value == "TOTAL"

    # Row 4: DATE header and day numbers
    assert ws.cell(row=4, column=3).value == "DATE"
    assert ws.cell(row=4, column=4).value == 3
    assert ws.cell(row=4, column=5).value == 7
    assert ws.cell(row=4, column=6).value == 10
    assert ws.cell(row=4, column=7).value == 3  # 3 dates * 1 hour = 3 total

    # Check student rows
    # Student 1: Arun (Roll 01) present on all 3 days -> 3 hrs, 100%, 5 marks
    assert ws.cell(row=8, column=2).value == "101"
    assert ws.cell(row=8, column=7).value == 3
    assert ws.cell(row=8, column=8).value == 100  # rounded off
    assert ws.cell(row=8, column=9).value == 5

    # Student 2: Bala (Roll 02) absent on 2 days, present on 1 day -> 1 hr, 33% (rounded off), 5 marks (2 absent < 3)
    assert ws.cell(row=9, column=2).value == "102"
    assert ws.cell(row=9, column=7).value == 1
    assert ws.cell(row=9, column=8).value == 33   # rounded off from 33.33%
    assert ws.cell(row=9, column=9).value == 5


def test_full_register_workflow_and_export():
    # 1. Configure session: Jan 5 to Jan 22, Monday & Thursday (0 and 3), 1 period (hour) per day
    start_d = date(2026, 1, 5)
    end_d = date(2026, 1, 22)
    save_register_session(start_d, end_d, [0, 3], periods_per_day=1, weekday_periods={0: 1, 3: 4})

    session = get_register_session()
    assert session is not None
    assert session["start_date"] == start_d
    assert session["end_date"] == end_d
    assert session["periods_per_day"] == 1
    assert session["weekday_periods"] == {0: 1, 3: 4}

    class_dates = get_session_class_dates(start_d, end_d, session["weekdays"])
    # Jan 2026: 5 (Mon), 8 (Thu), 12 (Mon), 15 (Thu), 19 (Mon), 22 (Thu) -> 6 class dates
    assert len(class_dates) == 6

    # 2. Initially, no dates are marked
    status = get_session_status(class_dates)
    assert status["is_complete"] is False
    assert status["total_dates"] == 6
    assert len(status["marked_dates"]) == 0
    assert status["next_unmarked_date"] == class_dates[0]

    # 3. Mark dates up to the second-to-last date
    for d in class_dates[:-1]:
        save_attendance(d, "02,04")

    status = get_session_status(class_dates)
    assert status["is_complete"] is False
    assert len(status["unmarked_dates"]) == 1
    assert status["next_unmarked_date"] == class_dates[-1]

    # 4. Mark the final date (Jan 22)
    save_attendance(class_dates[-1], "02")
    status = get_session_status(class_dates)
    assert status["is_complete"] is True
    assert status["next_unmarked_date"] is None

    # 5. Generate the register Excel
    excel_bytes = generate_register_excel(start_d, end_d, class_dates=class_dates, periods_per_day=1)
    wb = openpyxl.load_workbook(BytesIO(excel_bytes))
    ws = wb.active
    assert ws.title == "Attendance Register"

    # Total periods: 6 dates * 1 = 6 hours
    assert ws.cell(row=3, column=3).value == "PERIODS"
    assert ws.cell(row=3, column=4).value == 1  # Jan 5 (Mon) -> 1st hour
    assert ws.cell(row=3, column=5).value == 4  # Jan 8 (Thu) -> 4th hour
    assert ws.cell(row=4, column=10).value == 6  # TOTAL under Col J (col 10)
    assert ws.cell(row=5, column=4).value == "JAN"  # Merged Month

    # Check student 1 (Roll 01, Arun): 100% present -> 6 hrs, 100% (rounded), 5 marks
    assert ws.cell(row=8, column=2).value == "101"
    assert ws.cell(row=8, column=10).value == 6
    assert ws.cell(row=8, column=11).value == 100
    assert ws.cell(row=8, column=12).value == 5

    # Check student 2 (Roll 02, Bala): absent on all 6 dates -> 6 absent -> 5 - (6//3) = 3 marks
    assert ws.cell(row=9, column=2).value == "102"
    assert ws.cell(row=9, column=10).value == 0
    assert ws.cell(row=9, column=11).value == 0
    assert ws.cell(row=9, column=12).value == 3


def test_mark_holiday_by_input_and_session_status():
    d = date(2026, 9, 3)
    class_dates = [d]
    # Inputting 'holiday' marks date as holiday
    pres, ab = save_attendance(d, "holiday")
    assert (pres, ab) == (0, 0)
    # Check that session status marks it as completed
    status = get_session_status(class_dates)
    assert status["is_complete"] is True
    assert d in status["marked_dates"]
    assert status["unmarked_dates"] == []


def test_40_days_40_hours_with_holidays_and_period_selection():
    # Generate 40 Mondays and Thursdays: 20 weeks * 2 = 40 days
    # Starting from Jan 5, 2026
    start_d = date(2026, 1, 5)
    end_d = date(2026, 5, 21)
    class_dates = get_session_class_dates(start_d, end_d, weekdays=[0, 3])
    assert len(class_dates) == 40  # Exactly 40 class days

    save_register_session(start_d, end_d, [0, 3], periods_per_day=1, weekday_periods={0: 1, 3: 4})

    # Mark 2 days as holidays (e.g. date 5 and date 10) using 'holiday' text input
    hol_dates = {class_dates[4], class_dates[9]}
    for d in class_dates:
        if d in hol_dates:
            save_attendance(d, "holiday")
        else:
            save_attendance(d, "02")  # Bala (02) is absent, everyone else present

    # Check session completion status
    status = get_session_status(class_dates)
    assert status["is_complete"] is True
    assert len(status["marked_dates"]) == 40

    # Generate Excel register (by default include_holidays=False)
    excel_bytes = generate_register_excel(start_d, end_d, class_dates=class_dates, periods_per_day=1)
    wb = openpyxl.load_workbook(BytesIO(excel_bytes))
    ws = wb.active

    # Total working hours: 40 days - 2 holidays = 38 working days exported as columns!
    hrs_col = 4 + 38
    pct_col = 5 + 38
    marks_col = 6 + 38
    assert ws.cell(row=4, column=hrs_col).value == 38

    # Verify that NO cell contains "H" anywhere in the exported rows
    for r in range(1, 10):
        for c in range(1, marks_col + 1):
            assert ws.cell(row=r, column=c).value != "H"

    # Check student 1 (Roll 01, Arun): present on all 38 working days -> 38/38 = 100% -> 5 marks
    assert ws.cell(row=8, column=hrs_col).value == 38
    assert ws.cell(row=8, column=pct_col).value == 100
    assert ws.cell(row=8, column=marks_col).value == 5

    # Check student 2 (Roll 02, Bala): absent on all 38 working days -> 0/38 = 0% -> 0 marks
    assert ws.cell(row=9, column=hrs_col).value == 0
    assert ws.cell(row=9, column=pct_col).value == 0
    assert ws.cell(row=9, column=marks_col).value == 0

    # Also verify legacy include_holidays=True if explicitly requested
    legacy_bytes = generate_register_excel(start_d, end_d, class_dates=class_dates, periods_per_day=1, include_holidays=True)
    ws_leg = openpyxl.load_workbook(BytesIO(legacy_bytes)).active
    assert ws_leg.cell(row=4, column=4 + 40).value == 38


def test_student_roll_wise_natural_sorting_and_excel_arrangement():
    # Clear existing and add students with mixed string roll numbers: 1, 10, 2, 09, 11, 3
    database.init_db()
    from database import connection
    with connection() as conn:
        conn.execute("DELETE FROM daily_attendance")
        conn.execute("DELETE FROM students")

    test_data = [
        ("REG10", "10", "Ten"),
        ("REG01", "1", "One"),
        ("REG02", "2", "Two"),
        ("REG20", "20", "Twenty"),
        ("REG09", "09", "Nine"),
        ("REG11", "11", "Eleven"),
        ("REG03", "3", "Three"),
    ]
    for reg, roll, name in test_data:
        add_student(reg, roll, name)

    students = database.get_students() if hasattr(database, "get_students") else get_students()
    roll_order = [s["roll_no"] for s in students]
    # Must be naturally sorted: 1, 2, 3, 09, 10, 11, 20 (not 1, 10, 11, 2...)
    assert roll_order == ["1", "2", "3", "09", "10", "11", "20"]

    # In Excel register, Row 8 must be Roll 1, Row 9 Roll 2, Row 10 Roll 3, etc.
    d = date(2026, 9, 3)
    save_attendance(d, "2,11")  # Mark Roll 2 and Roll 11 absent
    excel_bytes = generate_register_excel(d, d, class_dates=[d], periods_per_day=1)
    wb = openpyxl.load_workbook(BytesIO(excel_bytes))
    ws = wb.active

    # Check student rows in Excel
    expected_names = ["One", "Two", "Three", "Nine", "Ten", "Eleven", "Twenty"]
    for idx, expected_name in enumerate(expected_names, start=1):
        row_num = 7 + idx
        assert ws.cell(row=row_num, column=1).value == idx  # SN: 1, 2, 3, 4, 5, 6, 7
        assert ws.cell(row=row_num, column=3).value == expected_name


def test_flexible_roll_number_matching_and_p_keyword():
    d = date(2026, 9, 14)
    # The default temporary_database has students with roll numbers "01", "02", "03", "04", "05"
    # Typing "1" should match "01", typing "2" should match "02"
    present, absent = save_attendance(d, "1, 2", replace=True)
    assert (present, absent) == (3, 2)
    assert get_absent_roll_nos(d) == ["01", "02"]

    # Typing "P" or "all" should mark all present
    present, absent = save_attendance(d, "P", replace=True)
    assert (present, absent) == (5, 0)
    assert get_absent_roll_nos(d) == []


def test_multiple_datasets_isolation_and_deletion(tmp_path, monkeypatch):
    test_datasets_dir = tmp_path / "datasets"
    monkeypatch.setattr(database, "DATASETS_DIR", test_datasets_dir)
    monkeypatch.setattr(database, "DB_PATH", database._default_db_path)
    monkeypatch.setattr(database, "_active_dataset", "Class A")

    # Create dataset Class A
    ds_a = database.create_dataset("Class A")
    assert ds_a == "Class A"
    assert database.get_active_dataset() == "Class A"
    database.init_db()
    add_student("REG_A1", "01", "Student In A")

    # Create dataset Class B
    ds_b = database.create_dataset("Class B")
    assert ds_b == "Class B"
    assert database.get_active_dataset() == "Class B"
    database.init_db()
    add_student("REG_B1", "01", "Student In B")  # Same roll number 01, completely isolated!

    # Verify Class B has only Student In B
    students_b = get_students()
    assert len(students_b) == 1
    assert students_b[0]["name"] == "Student In B"

    # Switch back to Class A
    database.set_active_dataset("Class A")
    students_a = get_students()
    assert len(students_a) == 1
    assert students_a[0]["name"] == "Student In A"

    # Verify list_datasets
    all_ds = database.list_datasets()
    assert "Class A" in all_ds and "Class B" in all_ds

    # Verify get_dataset_info
    info_a = database.get_dataset_info("Class A")
    assert info_a["student_count"] == 1

    # Delete Class A
    deleted = database.delete_dataset("Class A")
    assert deleted is True
    remaining = database.list_datasets()
    assert "Class A" not in remaining
    assert "Class B" in remaining
    assert database.get_active_dataset() == "Class B"
    assert get_students()[0]["name"] == "Student In B"


def test_course_metadata_storage_and_excel_header():
    save_course_details(
        class_handled="III CSE - A",
        faculty_name="Dr. K. Sharma",
        subject="Database Management Systems",
        total_hours=45,
    )
    info = get_course_details()
    assert info["class_handled"] == "III CSE - A"
    assert info["faculty_name"] == "Dr. K. Sharma"
    assert info["subject"] == "Database Management Systems"
    assert info["total_hours"] == 45

    # Generate Excel and inspect Row 1 and Row 2
    d = date(2026, 9, 14)
    save_attendance(d, "02", replace=True)
    excel_bytes = generate_register_excel(d, d, class_dates=[d], periods_per_day=1)
    wb = openpyxl.load_workbook(BytesIO(excel_bytes))
    ws = wb.active

    # Check Row 1
    assert ws.cell(row=1, column=1).value == "CLASS: III CSE - A"
    assert ws.cell(row=1, column=3).value == "SUBJECT: Database Management Systems"
    assert ws.cell(row=1, column=4).value == "FACULTY: Dr. K. Sharma"

    # Check Row 2
    assert ws.cell(row=2, column=1).value == "TOTAL HOURS: 45"

    # Verify standard register rows remain intact
    assert ws.cell(row=3, column=3).value == "PERIODS"
    assert ws.cell(row=4, column=3).value == "DATE"
    assert ws.cell(row=7, column=1).value == "SN"
    assert ws.cell(row=8, column=1).value == 1
    assert ws.cell(row=8, column=2).value == "101"
    assert ws.cell(row=8, column=3).value == "Arun"


def test_dataset_creation_duplicate_prevention_and_exists(tmp_path, monkeypatch):
    monkeypatch.setattr(database, "DATASETS_DIR", Path(tmp_path) / "datasets")
    monkeypatch.setattr(database, "DATA_DIR", Path(tmp_path))
    monkeypatch.setattr(database, "DB_PATH", database._default_db_path)

    # Initial state
    assert not database.dataset_exists("Math 101")
    database.create_dataset("Math 101")
    assert database.dataset_exists("Math 101")
    assert database.dataset_exists("math 101")  # case-insensitive

    # Creating with same name should raise ValueError unless auto_increment=True
    with pytest.raises(ValueError, match="already exists"):
        database.create_dataset("Math 101")

    # With auto_increment=True it creates Math 101 (2)
    new_name = database.create_dataset("Math 101", auto_increment=True)
    assert new_name == "Math 101 (2)"
    assert database.dataset_exists("Math 101 (2)")


def test_dataset_deletion_and_reset(tmp_path, monkeypatch):
    datasets_dir = Path(tmp_path) / "datasets"
    monkeypatch.setattr(database, "DATASETS_DIR", datasets_dir)
    monkeypatch.setattr(database, "DATA_DIR", Path(tmp_path))
    monkeypatch.setattr(database, "DB_PATH", database._default_db_path)

    database.create_dataset("Class X")
    database.create_dataset("Class Y")
    assert database.dataset_exists("Class X")
    assert database.dataset_exists("Class Y")

    # Delete Class X
    assert database.delete_dataset("Class X") is True
    assert not database.dataset_exists("Class X")
    assert database.dataset_exists("Class Y")

    # Delete Class Y (the last dataset) -> should reset to Class 1
    assert database.delete_dataset("Class Y") is True
    assert not database.dataset_exists("Class Y")
    assert database.dataset_exists("Class 1")


def test_lab_attendance_marks_scheme():
    # Lab Rule: 1 mark reduced for every 1 day absent (from max 5 marks)
    assert calculate_attendance_marks(0, total_periods=10, course_type="lab") == 5
    assert calculate_attendance_marks(1, total_periods=10, course_type="lab") == 4
    assert calculate_attendance_marks(2, total_periods=10, course_type="lab") == 3
    assert calculate_attendance_marks(3, total_periods=10, course_type="lab") == 2
    assert calculate_attendance_marks(4, total_periods=10, course_type="lab") == 1
    assert calculate_attendance_marks(5, total_periods=10, course_type="lab") == 0
    assert calculate_attendance_marks(6, total_periods=10, course_type="lab") == 0

    # Percentage conversions for lab
    assert calculate_attendance_marks(100.0, total_periods=10, course_type="lab") == 5
    assert calculate_attendance_marks(90.0, total_periods=10, course_type="lab") == 4  # 1 day absent
    assert calculate_attendance_marks(80.0, total_periods=10, course_type="lab") == 3  # 2 days absent
    assert calculate_attendance_marks(50.0, total_periods=10, course_type="lab") == 0  # 5 days absent


def test_lab_excel_register_3_hours_per_day_and_marks():
    # 10 Wednesdays in Jan-March 2026: Jan 7, 14, 21, 28, Feb 4, 11, 18, 25, Mar 4, 11
    start_d = date(2026, 1, 7)
    end_d = date(2026, 3, 11)
    class_dates = get_session_class_dates(start_d, end_d, weekdays=[2], exclude_holidays=False)
    assert len(class_dates) == 10

    # Mark attendance:
    # Student 1 (01 Arun): present all 10 days
    # Student 2 (02 Bala): absent on day 1 (1 day absent) -> 4 marks, 27 hrs
    # Student 3 (03 Deepa): absent on day 1 & 2 (2 days absent) -> 3 marks, 24 hrs
    # Student 4 (04 Hari): absent on 5 days -> 0 marks, 15 hrs
    save_attendance(class_dates[0], "02, 03, 04")
    save_attendance(class_dates[1], "03, 04")
    save_attendance(class_dates[2], "04")
    save_attendance(class_dates[3], "04")
    save_attendance(class_dates[4], "04")
    for d in class_dates[5:]:
        save_attendance(d, "P")

    excel_bytes = generate_register_excel(
        start_d,
        end_d,
        class_dates=class_dates,
        periods_per_day=3,
        weekday_periods={2: "1-3"},
        class_handled="III CSE Lab",
        subject="Database Lab",
        faculty_name="Dr. K. Sharma",
        total_hours=30,
        course_type="lab",
    )
    wb = openpyxl.load_workbook(BytesIO(excel_bytes))
    ws = wb.active
    assert ws.title == "Attendance Register"

    # Row 1 has Lab type
    assert ws.cell(row=1, column=5).value == "TYPE: LAB"

    # Row 3: PERIODS header and period string
    assert ws.cell(row=3, column=3).value == "PERIODS"
    assert ws.cell(row=3, column=4).value == "1-3"
    assert ws.cell(row=3, column=14).value == "TOTAL"

    # Row 4: Total working hours should be 10 dates * 3 hours = 30 hours!
    hrs_col = 4 + 10
    pct_col = 5 + 10
    marks_col = 6 + 10
    assert ws.cell(row=4, column=hrs_col).value == 30

    # Arun (01): 10 days present -> 30 hrs, 100%, 5 marks
    assert ws.cell(row=8, column=2).value == "101"
    assert ws.cell(row=8, column=hrs_col).value == 30
    assert ws.cell(row=8, column=pct_col).value == 100
    assert ws.cell(row=8, column=marks_col).value == 5

    # Bala (02): 1 day absent (9 days present) -> 27 hrs, 90%, 4 marks (1 mark reduced)
    assert ws.cell(row=9, column=2).value == "102"
    assert ws.cell(row=9, column=hrs_col).value == 27
    assert ws.cell(row=9, column=pct_col).value == 90
    assert ws.cell(row=9, column=marks_col).value == 4

    # Deepa (03): 2 days absent (8 days present) -> 24 hrs, 80%, 3 marks (2 marks reduced)
    assert ws.cell(row=10, column=2).value == "103"
    assert ws.cell(row=10, column=hrs_col).value == 24
    assert ws.cell(row=10, column=pct_col).value == 80
    assert ws.cell(row=10, column=marks_col).value == 3

    # Hari (04): 5 days absent (5 days present) -> 15 hrs, 50%, 0 marks (5 marks reduced)
    assert ws.cell(row=11, column=2).value == "104"
    assert ws.cell(row=11, column=hrs_col).value == 15
    assert ws.cell(row=11, column=pct_col).value == 50
    assert ws.cell(row=11, column=marks_col).value == 0


def test_holiday_editing_and_reopening():
    from attendance import delete_holiday, holiday_for_date
    d = date(2026, 9, 3)
    start_d = date(2026, 9, 1)
    end_d = date(2026, 9, 10)

    # 1. Mark as holiday
    save_attendance(d, "holiday")
    assert holiday_for_date(d) is not None

    # 2. Verify that with exclude_holidays=False, the date still appears in class_dates
    all_dates = get_session_class_dates(start_d, end_d, weekdays=[0, 3], exclude_holidays=False)
    assert d in all_dates

    # 3. User wants to edit/reopen the holiday: delete holiday
    delete_holiday(d)
    delete_attendance(d)
    assert holiday_for_date(d) is None
    assert not attendance_exists(d)

    # 4. User marks attendance directly over a holiday using replace=True
    save_attendance(d, "holiday")
    assert holiday_for_date(d) is not None
    # Now user changes mind and marks "P"
    save_attendance(d, "P", replace=True)
    assert holiday_for_date(d) is None
    assert attendance_exists(d)
    assert get_absent_roll_nos(d) == []


def test_course_details_storage_with_course_type():
    save_course_details("III CSE Lab", "Dr. Sharma", "Python Lab", total_hours=36, course_type="lab")
    details = get_course_details()
    assert details["course_type"] == "lab"
    assert details["class_handled"] == "III CSE Lab"
    assert details["total_hours"] == 36

    save_course_details("III CSE Theory", "Dr. Sharma", "Python Theory", total_hours=40, course_type="classroom")
    details2 = get_course_details()
    assert details2["course_type"] == "classroom"
    assert details2["total_hours"] == 40


def test_holidays_not_included_in_excel_export():
    import openpyxl
    start_d = date(2026, 8, 3)
    end_d = date(2026, 8, 31)
    # Mondays: Aug 3, 10, 17, 24, 31 -> 5 Mondays
    mondays = get_session_class_dates(start_d, end_d, weekdays=[0])
    assert len(mondays) == 5

    # Mark Aug 17 as a holiday
    save_attendance(date(2026, 8, 17), "holiday")
    # Mark the other 4 dates
    for d in [date(2026, 8, 3), date(2026, 8, 10), date(2026, 8, 24), date(2026, 8, 31)]:
        save_attendance(d, "")  # All present

    excel_bytes = generate_register_excel(start_d, end_d, class_dates=mondays, periods_per_day=3, course_type="lab")
    wb = openpyxl.load_workbook(BytesIO(excel_bytes))
    ws = wb.active

    # Exactly 4 dates should appear as columns (Aug 17 holiday must NOT appear)
    # Col 1: SN, Col 2: REGNO, Col 3: NAME, Cols 4-7: the 4 working dates, Col 8: TOTAL, Col 9: % ATT, Col 10: MARKS
    assert ws.cell(row=3, column=8).value == "TOTAL"
    assert ws.cell(row=4, column=8).value == 12  # 4 days * 3 hrs = 12 hrs

    # Verify dates in row 4: days should be 3, 10, 24, 31 (17 is absent from header)
    date_numbers = [ws.cell(row=4, column=col).value for col in range(4, 8)]
    assert date_numbers == [3, 10, 24, 31]
    assert 17 not in date_numbers

    # Verify that NO cell contains "H" anywhere in the sheet
    for r in range(1, 15):
        for c in range(1, 12):
            val = ws.cell(row=r, column=c).value
            assert val != "H", f"Found 'H' at row {r}, col {c}"

